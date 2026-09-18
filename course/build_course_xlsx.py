# -*- coding: utf-8 -*-
"""ガイド型（層2a）の Excel ブック course_moonbase.xlsx を組み立てる。

分析ステップ1〜5を、Excel の「見える数式」で行う。
- `course/build_course_data.py` が作った前処理済み CSV をシートに埋め込む
- 各ステップのシートには AVERAGEIFS / COUNTIFS / INDEX-MATCH などの数式と、
  グラフを置く（生徒は数値を書き換えて再計算する）
- ピボットテーブル・マクロ・Power Query は使わない（デスクトップ Excel の基本機能のみ）

    python course/build_course_data.py    # 先にこちら
    python course/build_course_xlsx.py

出力: course/course_moonbase.xlsx
"""
import pathlib

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows

HERE = pathlib.Path(__file__).resolve().parent
DATA = HERE / "data"
OUT = HERE / "course_moonbase.xlsx"

NAVY = "1A2F4B"
H1 = Font(name="Yu Gothic", size=15, bold=True, color="FFFFFF")
H2 = Font(name="Yu Gothic", size=12, bold=True, color=NAVY)
BODY = Font(name="Yu Gothic", size=10)
BOLD = Font(name="Yu Gothic", size=10, bold=True)
BLUE = Font(name="Yu Gothic", size=10, bold=True, color="1A6FB0")
HFILL = PatternFill("solid", fgColor=NAVY)
INFILL = PatternFill("solid", fgColor="FFF6E9")   # 生徒が入力するセル
WRAP = Alignment(wrap_text=True, vertical="top")


def _title(ws, text, row=1):
    ws.cell(row=row, column=1, value=text).font = H1
    ws.cell(row=row, column=1).fill = HFILL
    for c in range(2, 12):
        ws.cell(row=row, column=c).fill = HFILL
    ws.row_dimensions[row].height = 22


def _h2(ws, text, row):
    ws.cell(row=row, column=1, value=text).font = H2


def _note(ws, text, row, col=1, span=10):
    cell = ws.cell(row=row, column=col, value=text)
    cell.font = BODY
    cell.alignment = WRAP
    ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + span - 1)
    ws.row_dimensions[row].height = 14 * (1 + text.count("\n"))


def _input(ws, cell, value):
    ws[cell] = value
    ws[cell].fill = INFILL
    ws[cell].font = BOLD


def add_data_sheet(wb, csv_name, sheet_name):
    df = pd.read_csv(DATA / csv_name)
    ws = wb.create_sheet(sheet_name)
    for r in dataframe_to_rows(df, index=False, header=True):
        ws.append(r)
    for c in range(1, len(df.columns) + 1):
        ws.cell(row=1, column=c).font = BOLD
        ws.cell(row=1, column=c).fill = PatternFill("solid", fgColor="EEF1F6")
    ws.freeze_panes = "A2"
    ws.sheet_state = "visible"
    return ws, df


def build():
    wb = Workbook()
    wb.remove(wb.active)

    # ----- データシート（埋め込み） -----
    _, t_df = add_data_sheet(wb, "temp_grid.csv", "データ_温度")
    _, c_df = add_data_sheet(wb, "craters_labeled.csv", "データ_クレーター")
    _, a_df = add_data_sheet(wb, "crater_ages_labeled.csv", "データ_クレーター年代")
    ps_ws, ps_df = add_data_sheet(wb, "polar_south_sites.csv", "データ_南極")
    add_data_sheet(wb, "polar_north_sites.csv", "データ_北極")
    _, g_df = add_data_sheet(wb, "geology_grid.csv", "データ_地質")
    env_ws, env_df = add_data_sheet(wb, "env_grid.csv", "データ_環境")
    _, reg_df = add_data_sheet(wb, "candidate_regions.csv", "データ_地域")
    _, ls_df = add_data_sheet(wb, "landing_sites.csv", "データ_着陸地点")
    _, ref_df = add_data_sheet(wb, "reference.csv", "参考")

    T_N = len(t_df) + 1
    C_N = len(c_df) + 1
    A_N = len(a_df) + 1
    PS_N = len(ps_df) + 1
    G_N = len(g_df) + 1
    E_N = len(env_df) + 1
    REG_N = len(reg_df) + 1

    S3G = "ステップ3_月ぜんたい"
    S4R = "ステップ4_地域を選ぶ"

    # データ_環境 の列: A lat B lon C 区分 D age_index E temp_amp_K F night_min_K
    #   G noon_sun_elev_deg H earth_elev_deg I region
    #   J norm_sun_high K norm_amp_low L norm_earth_high M norm_earth_low N norm_night_warm → O スコア
    #   （全球傾斜 slope_deg は Python 版・data/site_environment.csv 側。表計算版は 5 指標）
    env_ws.cell(row=1, column=15, value="スコア").font = BOLD
    for r in range(2, E_N + 1):
        env_ws.cell(row=r, column=15, value=(
            f'=IF($I{r}={S4R}!$C$5,'
            f'({S4R}!$C$8*J{r}+{S4R}!$C$9*K{r}+{S4R}!$C$10*L{r}'
            f'+{S4R}!$C$11*M{r}+{S4R}!$C$12*N{r})'
            f'/MAX(1,{S4R}!$C$8+{S4R}!$C$9+{S4R}!$C$10+{S4R}!$C$11+{S4R}!$C$12),"")'))

    # データ_南極 の列: A lat B lon C illum D psf E km_to_shadow F slope_deg
    #   G norm_illum H norm_near_shadow I norm_low_psf J norm_low_slope K n  → L に「スコア」
    ps_ws.cell(row=1, column=12, value="スコア").font = BOLD
    S4B = "ステップ4b_スコア"
    for r in range(2, PS_N + 1):
        ps_ws.cell(row=r, column=12, value=(
            f"=({S4B}!$C$8*G{r}+{S4B}!$C$9*H{r}+{S4B}!$C$10*I{r}+{S4B}!$C$11*J{r})"
            f"/MAX(1,{S4B}!$C$8+{S4B}!$C$9+{S4B}!$C$10+{S4B}!$C$11)"))

    # データ_温度 の列: A=lat B=lon C..Z=t_lt00..23  AA=t_mean_K AB=t_swing_K
    T_SWING = f"データ_温度!$AB$2:$AB${T_N}"
    T_LAT = f"データ_温度!$A$2:$A${T_N}"

    # ================= はじめに =================
    ws = wb.create_sheet("はじめに", 0)
    _title(ws, "月データでムーンベースの場所を決めよう（Excel版）")
    _note(ws, "月の公開データ（温度・クレーター・極域の日照と傾斜・地質図）を Excel で分析して、"
              "「月面基地をどこに建てるか」を自分で決めます。コードは書きません。"
              "セルの数式を見て、黄色いセルの数字を書き換えて、結果をワークシートに記録します。", 3)
    _h2(ws, "1. ミッションを1つ選ぶ（ワークシートに○）", 5)
    ws["A6"], ws["B6"] = "ミッション", "基地に必要なこと"
    ws["A6"].font = ws["B6"].font = BOLD
    for i, (m, need) in enumerate([
        ("☀ 太陽光発電基地", "よく日が当たること。地面が平らなこと"),
        ("🔭 電波天文台", "地球の電波が届かないこと（＝月の裏側）。温度が安定していること"),
        ("❄ 氷採掘基地", "氷がありそうなこと（ずっと日が当たらない永久影のそば）。平らなこと"),
        ("🏠 有人総合基地", "電力・温度・氷・地球との通信をバランスさせること"),
    ]):
        ws.cell(row=7 + i, column=1, value=m).font = BODY
        ws.cell(row=7 + i, column=2, value=need).font = BODY
    _note(ws, "同じデータでも、ミッションが変われば「最適な場所」は変わります。"
              "氷採掘は南極に、電波天文台は裏側に、通信重視なら表側に――行き先は半球ごと変わります。", 12)

    _h2(ws, "2. 進め方（各ステップ共通）", 13)
    _note(ws, "① まずワークシートに『予想』を書く（数式を見る前に）\n"
              "② 黄色いセルの数字を書き換えて、結果（青いセル）を読む\n"
              "③ ワークシートに結果と『気づいたこと』を書く", 14)

    _h2(ws, "3. シートの並び", 18)
    for i, (nm, desc) in enumerate([
        ("ステップ1_温度", "月の1日の温度は緯度でどう変わる？ どこまで信じられる？"),
        ("ステップ2_海と陸", "『海』と『陸』でクレーターの数・大きさ・年代はどう違う？ 地質図と合っている？"),
        ("ステップ3_月ぜんたい", "月ぜんたいで環境を見る（日較差・地球の仰角・太陽高度）。どこも『全部で一番』にはならない"),
        ("ステップ4_地域を選ぶ", "ミッションに合う地域タイプを選び、その中でいちばんよい場所を点数で決める"),
        ("ステップ4b_南極", "（氷採掘・南極を選んだ班）南極の日照・傾斜・永久影を細かく見る"),
        ("ステップ5_まとめ", "各ミッションの答え（南極・裏側・表側…）と実在の計画を見比べる"),
        ("データ_◯◯ / 参考", "分析のもとデータ（前処理済み）。作り方は course/build_course_data.py"),
    ]):
        ws.cell(row=19 + i, column=1, value=nm).font = BOLD
        ws.cell(row=19 + i, column=2, value=desc).font = BODY
    _note(ws, "参考：同じ分析を Python（pandas）で書くと？ は notebooks/course_moonbase.ipynb にあります。"
              "Excel でやったことと数値がそろうように作ってあります。", 27)
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 62

    # ================= ステップ1：温度 =================
    ws = wb.create_sheet("ステップ1_温度")
    _title(ws, "ステップ1：月の温度は1日でどれくらい変わる？")
    _note(ws, "月には大気がない＝温室効果も熱の運搬もない。太陽が当たる昼と、当たらない夜の差が大きい。"
              "「1日の温度の較差（＝最高－最低）」を、緯度帯ごとに平均して比べる。"
              "データ『データ_温度』は3度ごとの世界地図（7200地点）。各地点に t_swing_K（1日の較差）が入っている。", 3)

    _h2(ws, "A. 緯度帯ごとの1日の温度較差（黄色いセルに緯度を入れる）", 6)
    ws["A7"], ws["B7"], ws["C7"], ws["D7"] = "緯度の下", "緯度の上", "地点数", "1日の較差の平均 [K]"
    for c in "ABCD":
        ws[c + "7"].font = BOLD
    bands = [(-6, 6, "赤道"), (24, 36, "中緯度"), (54, 66, "高緯度"), (78, 90, "極付近（要注意）")]
    for i, (lo, hi, label) in enumerate(bands):
        r = 8 + i
        _input(ws, f"A{r}", lo)
        _input(ws, f"B{r}", hi)
        ws[f"C{r}"] = f'=COUNTIFS({T_LAT},">="&A{r},{T_LAT},"<="&B{r})'
        ws[f"D{r}"] = (f'=ROUND(AVERAGEIFS({T_SWING},{T_LAT},">="&A{r},'
                       f'{T_LAT},"<="&B{r}),0)')
        ws[f"D{r}"].font = BLUE
        ws[f"E{r}"] = label
        ws[f"E{r}"].font = BODY
    _note(ws, "気づき：緯度が高くなると較差は？　赤道の較差（約○○K＝約○○℃）を、"
              "地球の砂漠の昼夜差（20〜30℃）と比べると？　極付近の値は信じてよい？（下の C も見る）", 13)

    _h2(ws, "B. 1地点の24時間カーブを見る（黄色いセルに地点を入れる）", 16)
    _input(ws, "B17", 0.25)
    _input(ws, "B18", 0.25)
    ws["A17"], ws["A18"] = "見たい地点の緯度", "見たい地点の経度"
    ws["A17"].font = ws["A18"].font = BODY
    _note(ws, "※ 緯度・経度は『データ_温度』にある値ちょうどを入れる（緯度は …-2.75, 0.25, 3.25, 6.25…／"
              "経度も同じきざみ）。見つからないと温度が 0 のまま。", 19, span=8)
    ws["A21"] = "現地時間"
    ws["B21"] = "温度 [K]"
    ws["A21"].font = ws["B21"].font = BOLD
    for h in range(24):
        r = 22 + h
        ws[f"A{r}"] = h
        col = get_column_letter(3 + h)   # t_lt00 は データ_温度 の C列
        ws[f"B{r}"] = (f'=SUMIFS(データ_温度!${col}$2:${col}${T_N},'
                       f'データ_温度!$A$2:$A${T_N},$B$17,'
                       f'データ_温度!$B$2:$B${T_N},$B$18)')
    line = LineChart()
    line.title = "選んだ地点の1日の温度カーブ"
    line.y_axis.title = "温度 [K]"
    line.x_axis.title = "現地時間"
    data = Reference(ws, min_col=2, min_row=21, max_row=45)
    cats = Reference(ws, min_col=1, min_row=22, max_row=45)
    line.add_data(data, titles_from_data=True)
    line.set_categories(cats)
    line.height, line.width = 8, 15
    ws.add_chart(line, "D16")
    _note(ws, "気づき：いちばん暑い時刻・いちばん寒い時刻はいつ？　朝と夕方でカーブの形は左右対称？　"
              "緯度を -86.75（極付近）にすると、カーブはどうなる？　このデータはそこで信じてよい？", 47)

    _h2(ws, "C. 実際に人が降りた場所の温度（『データ_着陸地点』より）", 50)
    ws["A51"], ws["B51"], ws["C51"], ws["D51"] = "着陸地点", "正午 [K]", "真夜中 [K]", "1日の差 [K]"
    for c in "ABCD":
        ws[c + "51"].font = BOLD
    for i, nm in enumerate(["Apollo 11", "Apollo 15", "Chang'e 4", "Chandrayaan-3"]):
        r = 52 + i
        ws[f"A{r}"] = nm
        ws[f"A{r}"].font = BODY
        m = f'MATCH("{nm}",データ_着陸地点!$A$2:$A${len(ls_df) + 1},0)'
        ws[f"B{r}"] = f'=INDEX(データ_着陸地点!$G$2:$G${len(ls_df) + 1},{m})'
        ws[f"C{r}"] = f'=INDEX(データ_着陸地点!$H$2:$H${len(ls_df) + 1},{m})'
        ws[f"D{r}"] = f'=INDEX(データ_着陸地点!$I$2:$I${len(ls_df) + 1},{m})'
        for c in "BCD":
            ws[f"{c}{r}"].font = BLUE
    _note(ws, "気づき：赤道の海（Apollo 11）と高緯度（Chandrayaan-3）で、1日の差はどう違う？", 56)
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 16
    for c in "CDE":
        ws.column_dimensions[c].width = 13

    # ================= ステップ2：海と陸 =================
    ws = wb.create_sheet("ステップ2_海と陸")
    _title(ws, "ステップ2：月の「海」と「陸」で何が違う？")
    _note(ws, "月の黒い部分＝『海』（マリア）は、昔の溶岩でおおわれた低地。"
              "『データ_クレーター』の各行に、その場所が海か陸かの区分がついている（円で囲む近似。作り方は『参考』シート）。"
              "数（密度）・大きさ・年代を海と陸で比べ、最後に USGS の公式地質図と『答え合わせ』する。", 3)
    C_KUBUN = f"データ_クレーター!$D$2:$D${C_N}"
    C_DIAM = f"データ_クレーター!$C$2:$C${C_N}"
    A_KUBUN = f"データ_クレーター年代!$F$2:$F${A_N}"
    A_AGE = f"データ_クレーター年代!$D$2:$D${A_N}"
    G_KUBUN = f"データ_地質!$E$2:$E${G_N}"
    G_AGE = f"データ_地質!$D$2:$D${G_N}"

    _h2(ws, "A. クレーターの数と密度", 6)
    rows = [
        ("海のクレーター数", f'=COUNTIF({C_KUBUN},"海")'),
        ("陸のクレーター数", f'=COUNTIF({C_KUBUN},"陸")'),
        ("海の面積割合（円近似・参考シート）", "=参考!B2"),
        ("陸の面積割合（円近似）", "=参考!B3"),
        ("海の密度（数÷面積割合）", "=ROUND(B7/B9)"),
        ("陸の密度（数÷面積割合）", "=ROUND(B8/B10)"),
        ("陸は海の何倍こみあっている？", "=ROUND(B12/B11,1)"),
    ]
    for i, (label, formula) in enumerate(rows):
        r = 7 + i
        ws[f"A{r}"] = label
        ws[f"A{r}"].font = BODY
        ws[f"B{r}"] = formula
        ws[f"B{r}"].font = BLUE

    _h2(ws, "B. クレーターの大きさ（直径の平均）", 16)
    ws["A17"], ws["B17"] = "海の直径の平均 [km]", f'=ROUND(AVERAGEIF({C_KUBUN},"海",{C_DIAM}),1)'
    ws["A18"], ws["B18"] = "陸の直径の平均 [km]", f'=ROUND(AVERAGEIF({C_KUBUN},"陸",{C_DIAM}),1)'

    _h2(ws, "C. クレーターの年代（1=最も古い 〜 5=最も新しい）", 20)
    ws["A21"], ws["B21"] = "海の年代の平均", f'=ROUND(AVERAGEIF({A_KUBUN},"海",{A_AGE}),2)'
    ws["A22"], ws["B22"] = "陸の年代の平均", f'=ROUND(AVERAGEIF({A_KUBUN},"陸",{A_AGE}),2)'
    for r in (17, 18, 21, 22):
        ws[f"A{r}"].font = BODY
        ws[f"B{r}"].font = BLUE

    _h2(ws, "D. 答え合わせ：USGS の公式地質図と比べる", 25)
    checks = [
        ("あなたの推定：海と陸、古いのはどっち？", "（C の年代の平均が小さいほう）", None),
        ("地質図：海の相対年代の平均", "＝参考!B6", "=参考!B6"),
        ("地質図：陸の相対年代の平均", "＝参考!B7（大きいほど新しい）", "=参考!B7"),
        ("海の面積割合：あなたの円近似", "＝参考!B2", "=参考!B2"),
        ("海の面積割合：USGS 地質図", "＝参考!B4（文献値 約16%）", "=参考!B4"),
        ("円近似は USGS より大きい？小さい？その差は何%ポイント？", "=ROUND((参考!B2-参考!B4)*100,1)", "=ROUND((参考!B2-参考!B4)*100,1)"),
    ]
    for i, (label, hint, formula) in enumerate(checks):
        r = 26 + i
        ws[f"A{r}"] = label
        ws[f"A{r}"].font = BODY
        if formula:
            ws[f"B{r}"] = formula
            ws[f"B{r}"].font = BLUE
        ws[f"C{r}"] = hint
        ws[f"C{r}"].font = Font(name="Yu Gothic", size=9, color="777777")
    _note(ws, "気づき：①海のほうがクレーターが少ないのはなぜ？（隕石が落ちなかった？ 落ちた後に消えた？）"
              "②あなたのクレーターからの推定は、地質図と合っていた？　"
              "③円で海を囲むと面積割合が大きく出るのはなぜ？（海岸線の外側の陸も丸に入るから）", 33)
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 30

    # ================= ステップ3：月ぜんたいで環境を見る =================
    ws = wb.create_sheet(S3G)
    _title(ws, "ステップ3：月ぜんたいで環境を見る")
    _note(ws, "基地の場所を決めるには、温度だけでなく『地球が見えるか（通信・電波静穏）』『太陽がどれだけ高いか』"
              "も要る。『データ_環境』は月ぜんたいを3度マスに区切った表。temp_amp_K=1日の温度差、"
              "night_min_K=夜の最低温度、noon_sun_elev_deg=正午の太陽高度、"
              "earth_elev_deg=地球の仰角（正＝表側で通信できる／負＝裏側で電波が静か）。", 3)
    E_LAT = f"データ_環境!$A$2:$A${E_N}"
    E_AMP = f"データ_環境!$E$2:$E${E_N}"
    E_EARTH = f"データ_環境!$H$2:$H${E_N}"
    E_SUN = f"データ_環境!$G$2:$G${E_N}"
    E_REG = f"データ_環境!$I$2:$I${E_N}"
    E_NIGHT = f"データ_環境!$F$2:$F${E_N}"

    _h2(ws, "A. 緯度帯ごとの1日の温度差（黄色いセルに緯度）", 6)
    ws["A7"], ws["B7"], ws["C7"], ws["D7"] = "緯度の下", "緯度の上", "地点数", "1日の温度差の平均 [K]"
    for c in "ABCD":
        ws[c + "7"].font = BOLD
    for i, (lo, hi, label) in enumerate([(-6, 6, "赤道"), (24, 36, "中緯度"),
                                         (54, 66, "高緯度"), (81, 90, "極付近")]):
        r = 8 + i
        _input(ws, f"A{r}", lo)
        _input(ws, f"B{r}", hi)
        ws[f"C{r}"] = f'=COUNTIFS({E_LAT},">="&A{r},{E_LAT},"<="&B{r})'
        ws[f"D{r}"] = f'=ROUND(AVERAGEIFS({E_AMP},{E_LAT},">="&A{r},{E_LAT},"<="&B{r}),0)'
        ws[f"D{r}"].font = BLUE
        ws[f"E{r}"] = label
        ws[f"E{r}"].font = BODY

    _h2(ws, "B. 地域タイプごとの環境（『データ_地域』の名前で引く）", 14)
    ws["A15"], ws["B15"], ws["C15"], ws["D15"], ws["E15"] = (
        "地域タイプ", "地点数", "日較差の平均 [K]", "太陽高度の平均 [度]", "地球の仰角の平均 [度]")
    for c in "ABCDE":
        ws[c + "15"].font = BOLD
    for i, nm in enumerate(["赤道の海（静かの海）", "中緯度の火砕丘（Aristarchus 高原）",
                            "裏側・赤道（電波天文の候補域）", "南極（Shackleton-de Gerlache）"]):
        r = 16 + i
        ws[f"A{r}"] = nm
        ws[f"A{r}"].font = BODY
        ws[f"B{r}"] = f'=COUNTIF({E_REG},$A{r})'
        ws[f"C{r}"] = f'=ROUND(AVERAGEIF({E_REG},$A{r},{E_AMP}),0)'
        ws[f"D{r}"] = f'=ROUND(AVERAGEIF({E_REG},$A{r},{E_SUN}),0)'
        ws[f"E{r}"] = f'=ROUND(AVERAGEIF({E_REG},$A{r},{E_EARTH}),0)'
        for c in "BCDE":
            ws[f"{c}{r}"].font = BLUE
    _note(ws, "気づき：『温度が安定』『地球が見える』『日がよく当たる』が全部そろう地域はあった？　"
              "南極は日較差が小さいが太陽高度は？　裏側は地球の仰角が負（＝地球が地平線の下）。"
              "あなたのミッションで、いちばん大事な列はどれ？", 22)
    ws.column_dimensions["A"].width = 32
    for c in "BCDE":
        ws.column_dimensions[c].width = 16

    # ================= ステップ4：地域を選んで評価する =================
    ws = wb.create_sheet(S4R)
    _title(ws, "ステップ4：地域タイプを選んで、ミッションに合わせて評価する")
    _note(ws, "『データ_地域』からミッションに合う地域タイプを1つ選び（C5 に名前をそのまま書く）、"
              "5つの指標を 0〜1 に直して重みをつけて合計する（＝あなたのスコア式）。"
              "『データ_環境』の O 列『スコア』が、選んだ地域の行だけ自動で計算される。", 3)
    ws["A5"] = "選んだ地域タイプ（『データ_地域』の name をそのまま）"
    ws["A5"].font = BODY
    _input(ws, "C5", "赤道の海（静かの海）")

    _h2(ws, "A. あなたのスコア式：重み（黄色いセル。0なら気にしない）", 7)
    for i, (label, good, w) in enumerate([
        ("太陽高度 noon_sun_elev_deg", "高いほどよい（発電）", 2),
        ("1日の温度差 temp_amp_K", "小さいほどよい（熱の安定）", 2),
        ("地球の仰角 earth_elev_deg（表側）", "高いほどよい（通信できる）", 0),
        ("地球の仰角 earth_elev_deg（裏側）", "低いほどよい（電波が静か・天文台）", 0),
        ("夜の最低温度 night_min_K", "高いほどよい（夜に冷えすぎない）", 0),
    ]):
        r = 8 + i
        ws[f"A{r}"] = label
        ws[f"A{r}"].font = BODY
        ws[f"B{r}"] = good
        ws[f"B{r}"].font = BODY
        _input(ws, f"C{r}", w)
    ws["A13"] = "重みの合計"
    ws["A13"].font = BOLD
    ws["B13"] = "=C8+C9+C10+C11+C12"
    ws["B13"].font = BLUE

    _h2(ws, "B. ミッションごとの選び方（一例。C5 と重みを入れ直して使う）", 15)
    ws["A16"], ws["B16"], ws["C16"] = "ミッション", "選ぶ地域タイプ", "重み（太陽・温度差・地球表・地球裏・夜）"
    for c in "ABC":
        ws[c + "16"].font = BOLD
    for i, (m, reg, w) in enumerate([
        ("☀ 太陽光発電", "赤道の海（静かの海） か 南極", "2・2・0・0・1"),
        ("🔭 電波天文台", "裏側・赤道（電波天文の候補域）", "0・2・0・3・0"),
        ("📡 通信中継", "赤道の海（静かの海）", "1・1・3・0・0"),
        ("🏠 有人総合", "いくつか試す（＋ステップ4b）", "1・2・1・0・1"),
        ("❄ 氷採掘", "南極 → ステップ4b へ（永久影は環境データに無い）", "―"),
    ]):
        r = 17 + i
        ws[f"A{r}"] = m
        ws[f"B{r}"] = reg
        ws[f"C{r}"] = w
        for c in "ABC":
            ws[f"{c}{r}"].font = BODY

    _h2(ws, "C. スコアの高い順トップ10（C5・重みを変えると入れ替わる）", 24)
    E_SCORE = f"データ_環境!$O$2:$O${E_N}"
    ws["A25"], ws["B25"], ws["C25"], ws["D25"], ws["E25"], ws["F25"], ws["G25"] = (
        "順位", "スコア", "緯度", "経度", "日較差 [K]", "太陽高度 [度]", "地球の仰角 [度]")
    for c in "ABCDEFG":
        ws[c + "25"].font = BOLD
    for k in range(1, 11):
        r = 25 + k
        ws[f"A{r}"] = k
        ws[f"B{r}"] = f'=IFERROR(ROUND(LARGE({E_SCORE},A{r}),3),"")'
        m = f'MATCH(LARGE({E_SCORE},A{r}),{E_SCORE},0)'
        ws[f"C{r}"] = f'=IFERROR(INDEX(データ_環境!$A$2:$A${E_N},{m}),"")'
        ws[f"D{r}"] = f'=IFERROR(INDEX(データ_環境!$B$2:$B${E_N},{m}),"")'
        ws[f"E{r}"] = f'=IFERROR(INDEX(データ_環境!$E$2:$E${E_N},{m}),"")'
        ws[f"F{r}"] = f'=IFERROR(INDEX(データ_環境!$G$2:$G${E_N},{m}),"")'
        ws[f"G{r}"] = f'=IFERROR(INDEX(データ_環境!$H$2:$H${E_N},{m}),"")'
        for c in "BCDEFG":
            ws[f"{c}{r}"].font = BLUE
    _note(ws, "気づき：選んだ地域のトップの場所は、どんな特徴？　"
              "☀太陽光は『赤道の海』と『南極』の両方で試して、明るさ（太陽高度）と熱の安定を比べる。"
              "🔭電波天文は『地球の仰角（裏側）』の重みを大きくすると、経度180度あたりが上位に来る。", 37)
    ws.column_dimensions["A"].width = 34
    for c in "BCDEFG":
        ws.column_dimensions[c].width = 14

    # ================= ステップ4b：南極の日照と傾斜 =================
    ws = wb.create_sheet("ステップ4b_南極")
    _title(ws, "ステップ4b：南極でよい場所（日当たり・平ら）はどれくらい？")
    _note(ws, "極では太陽が地平線近くを回るだけなので『昼夜』がない。地形の高いところは年中日が当たり、"
              "クレーターの底は年中影（永久影）。基地には日当たりだけでなく『地面が平ら（傾斜が小さい）』ことも要る。"
              "『データ_南極』の average_illumination_percent は年間日照率、slope_deg は傾斜[度]。"
              "※どちらも絶対値は他の資料と単純比較しない（「暗い/明るい」「平ら/急」の順番だけ信じる）。", 3)
    PS_ILLUM = f"データ_南極!$C$2:$C${PS_N}"
    PS_SLOPE = f"データ_南極!$F$2:$F${PS_N}"
    PS_PSF = f"データ_南極!$D$2:$D${PS_N}"

    _h2(ws, "A. しきい値より日照率が高い地点の数（黄色いセルにしきい値[%]）", 6)
    ws["A7"], ws["B7"], ws["C7"] = "しきい値 [%]", "その値以上の地点数", "全体に占める割合"
    for c in "ABC":
        ws[c + "7"].font = BOLD
    for i, th in enumerate([20, 30, 35, 40]):
        r = 8 + i
        _input(ws, f"A{r}", th)
        ws[f"B{r}"] = f'=COUNTIF({PS_ILLUM},">="&A{r})'
        ws[f"C{r}"] = f'=TEXT(B{r}/{PS_N - 1},"0.0%")'
        ws[f"B{r}"].font = BLUE

    _h2(ws, "B. 傾斜がしきい値より小さい（平らな）地点の数（黄色いセルに傾斜[度]）", 13)
    ws["A14"], ws["B14"], ws["C14"] = "傾斜 [度] 以下", "その地点数", "割合"
    for c in "ABC":
        ws[c + "14"].font = BOLD
    for i, th in enumerate([5, 8, 10, 15]):
        r = 15 + i
        _input(ws, f"A{r}", th)
        ws[f"B{r}"] = f'=COUNTIF({PS_SLOPE},"<="&A{r})'
        ws[f"C{r}"] = f'=TEXT(B{r}/{PS_N - 1},"0.0%")'
        ws[f"B{r}"].font = BLUE

    _h2(ws, "C. 「日当たりがよい」かつ「平ら」かつ「永久影のそば」は何地点？", 20)
    _input(ws, "B21", 30)
    _input(ws, "B22", 10)
    _input(ws, "B23", 20)
    ws["A21"], ws["A22"], ws["A23"] = "日照率 ≧ [%]", "傾斜 ≦ [度]", "永久影まで ≦ [km]"
    for r in (21, 22, 23):
        ws[f"A{r}"].font = BODY
    ws["A24"] = "3つとも満たす地点数"
    ws["A24"].font = BOLD
    ws["B24"] = (f'=COUNTIFS({PS_ILLUM},">="&B21,{PS_SLOPE},"<="&B22,'
                 f'データ_南極!$E$2:$E${PS_N},"<="&B23)')
    ws["B24"].font = BLUE

    # ヒストグラム（傾斜）
    _h2(ws, "D. 傾斜のヒストグラム（2度きざみ）", 27)
    ws["A28"], ws["B28"] = "階級（以上）", "地点数"
    ws["A28"].font = ws["B28"].font = BOLD
    for i, lo in enumerate(range(0, 30, 2)):
        r = 29 + i
        ws[f"A{r}"] = lo
        ws[f"B{r}"] = f'=COUNTIFS({PS_SLOPE},">="&A{r},{PS_SLOPE},"<"&(A{r}+2))'
    bar = BarChart()
    bar.title = "南極の傾斜の分布"
    bar.y_axis.title = "地点数"
    bar.x_axis.title = "傾斜 [度] 以上"
    bar.add_data(Reference(ws, min_col=2, min_row=28, max_row=43), titles_from_data=True)
    bar.set_categories(Reference(ws, min_col=1, min_row=29, max_row=43))
    bar.height, bar.width = 8, 14
    ws.add_chart(bar, "D27")
    _note(ws, "気づき：日当たりのよい地点と、平らな地点は『同じ場所』？　"
              "C で3つとも満たす地点はどれくらい残った？　"
              "いちばん平らな地点（傾斜が最小）を『データ_南極』で探すと、そこは日が当たる？", 46)
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 14

    # ================= ステップ4b：南極のスコア =================
    ws = wb.create_sheet(S4B)
    _title(ws, "ステップ4b：南極で、ミッションごとにいちばんよい場所を点数で決める")
    _note(ws, "（氷採掘・南極を選んだ班向け）『データ_南極』の4つの指標を 0〜1 の点数に直して、"
              "重みをつけて合計する（＝あなたのスコア式）。重みは黄色いセルで変える。", 3)
    _h2(ws, "A. あなたのスコア式：重み（黄色いセル。0なら「気にしない」）", 6)
    ws["A7"], ws["B7"], ws["C7"] = "指標", "よい向き", "重み"
    for c in "ABC":
        ws[c + "7"].font = BOLD
    weights = [
        ("日照率 (illum)", "高いほどよい", 3),
        ("永久影までの距離 (km_to_shadow)", "近いほどよい", 0),
        ("永久影率 (permanent_shadow_fraction)", "低いほどよい", 0),
        ("傾斜 (slope_deg)", "低いほどよい", 2),
    ]
    for i, (label, good, w) in enumerate(weights):
        r = 8 + i
        ws[f"A{r}"] = label
        ws[f"A{r}"].font = BODY
        ws[f"B{r}"] = good
        ws[f"B{r}"].font = BODY
        _input(ws, f"C{r}", w)
    ws["A12"] = "重みの合計"
    ws["A12"].font = BOLD
    ws["B12"] = "=C8+C9+C10+C11"
    ws["B12"].font = BLUE

    _h2(ws, "B. ミッションごとの重み（ワークシートに書く。上の黄色いセルに入れ直して使う）", 14)
    ws["A15"], ws["B15"], ws["C15"], ws["D15"], ws["E15"] = (
        "ミッション", "日照", "永久影まで", "永久影率", "傾斜")
    for c in "ABCDE":
        ws[c + "15"].font = BOLD
    for i, (m, w) in enumerate([
        ("☀ 太陽光発電基地", (3, 0, 0, 2)),
        ("❄ 氷採掘基地", (0, 3, 0, 2)),
        ("🏠 有人基地", (2, 2, 1, 2)),
    ]):
        r = 16 + i
        ws[f"A{r}"] = m
        ws[f"A{r}"].font = BODY
        for c, v in zip("BCDE", w):
            ws[f"{c}{r}"] = v
            ws[f"{c}{r}"].font = BODY
    _note(ws, "※これは一例。自分のミッションで「何を重く見るか」を考えて、A の黄色いセルに入れて使う。", 20)

    _note(ws, "『データ_南極』シートの右端（L 列）に『スコア』列があり、A の重みで自動計算される。"
              "0〜1の点数（norm_illum など）は前処理済み（作り方は『参考』と build_course_data.py）。", 22)

    _h2(ws, "C. スコアの高い順トップ10（重みを変えると入れ替わる）", 24)
    PS_SCORE = f"データ_南極!$L$2:$L${PS_N}"
    ws["A25"], ws["B25"], ws["C25"], ws["D25"], ws["E25"], ws["F25"], ws["G25"] = (
        "順位", "スコア", "緯度", "経度", "日照率 [%]", "傾斜 [度]", "永久影まで [km]")
    for c in "ABCDEFG":
        ws[c + "25"].font = BOLD
    for k in range(1, 11):
        r = 25 + k
        ws[f"A{r}"] = k
        ws[f"B{r}"] = f'=ROUND(LARGE({PS_SCORE},A{r}),3)'
        m = f'MATCH(LARGE({PS_SCORE},A{r}),{PS_SCORE},0)'
        ws[f"C{r}"] = f'=INDEX(データ_南極!$A$2:$A${PS_N},{m})'
        ws[f"D{r}"] = f'=INDEX(データ_南極!$B$2:$B${PS_N},{m})'
        ws[f"E{r}"] = f'=INDEX(データ_南極!$C$2:$C${PS_N},{m})'
        ws[f"F{r}"] = f'=INDEX(データ_南極!$F$2:$F${PS_N},{m})'
        ws[f"G{r}"] = f'=INDEX(データ_南極!$E$2:$E${PS_N},{m})'
        for c in "BCDEFG":
            ws[f"{c}{r}"].font = Font(name="Yu Gothic", size=10, color="1A6FB0")
    _note(ws, "気づき：あなたのミッションのトップの場所は、どんな特徴（日照・傾斜・永久影までの距離）？　"
              "日照を重くすると傾斜は？　両方を同時に満たす場所はあった？　重みを変えると順位はどう動く？", 37)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 14
    for c in "CDEFG":
        ws.column_dimensions[c].width = 13

    # ================= ステップ5：まとめ =================
    ws = wb.create_sheet("ステップ5_まとめ")
    _title(ws, "ステップ5：各ミッションの答えと実在の計画を見比べる")
    _note(ws, "各班が、選んだ地域タイプとトップの場所（緯度・経度）を下の表に書き写す。"
              "同じ月・同じデータなのに、ミッションによって答えは『半球ごと』変わったはず。", 3)
    ws["A6"], ws["B6"], ws["C6"], ws["D6"], ws["E6"] = (
        "ミッション", "選んだ地域タイプ", "トップの緯度", "トップの経度", "どの半球？")
    for c in "ABCDE":
        ws[c + "6"].font = BOLD
    for i, m in enumerate(["☀ 太陽光発電基地", "🔭 電波天文台", "❄ 氷採掘基地", "🏠 有人総合基地"]):
        r = 7 + i
        ws[f"A{r}"] = m
        ws[f"A{r}"].font = BODY
        for c in "BCDE":
            _input(ws, f"{c}{r}", "")

    _h2(ws, "B. 実在の計画は、ミッションで半球がちがう", 12)
    ws["A13"], ws["B13"], ws["C13"], ws["D13"] = "計画", "半球", "要件", "参考：地球の仰角"
    for c in "ABCD":
        ws[c + "13"].font = BOLD
    for i, (nm, where, req, ee) in enumerate([
        ("Artemis III（有人・氷）", "南極", "永久影の氷＋近くの日照尾根", "約 0°（地平線すれすれ）"),
        ("LCRT（月裏側電波望遠鏡・構想）", "裏側", "地球の電波が届かないこと", "負（地球が地平線の下）"),
        ("Apollo 11（赤道・実績）", "表側の海", "アボート容易・通信良好", "約 +66°（ほぼ真上）"),
        ("Chang'e 4/6（裏側・実績）", "裏側", "裏側の地質サンプル", "負"),
    ]):
        r = 14 + i
        ws[f"A{r}"] = nm
        ws[f"B{r}"] = where
        ws[f"C{r}"] = req
        ws[f"D{r}"] = ee
        for c in "ABCD":
            ws[f"{c}{r}"].font = BODY

    _h2(ws, "C. NASA Artemis III の南極候補地（『データ_着陸地点』より。氷・有人ミッション向け）", 20)
    ws["A21"], ws["B21"], ws["C21"], ws["D21"], ws["E21"] = (
        "候補地", "緯度", "経度", "傾斜 [度]", "日照率 [%]")
    for c in "ABCDE":
        ws[c + "21"].font = BOLD
    LSN = len(ls_df) + 1
    for i, nm in enumerate(["Artemis III: Malapert Massif", "Artemis III: Haworth",
                            "Artemis III: Nobile Rim 1", "Artemis III: de Gerlache Rim 2"]):
        r = 22 + i
        ws[f"A{r}"] = nm.replace("Artemis III: ", "")
        ws[f"A{r}"].font = BODY
        m = f'MATCH("{nm}",データ_着陸地点!$A$2:$A${LSN},0)'
        ws[f"B{r}"] = f'=INDEX(データ_着陸地点!$C$2:$C${LSN},{m})'
        ws[f"C{r}"] = f'=INDEX(データ_着陸地点!$D$2:$D${LSN},{m})'
        ws[f"D{r}"] = f'=INDEX(データ_着陸地点!$L$2:$L${LSN},{m})'
        ws[f"E{r}"] = f'=INDEX(データ_着陸地点!$M$2:$M${LSN},{m})'
        for c in "BCDE":
            ws[f"{c}{r}"].font = BLUE

    _note(ws, "考察：①なぜ班ごとに『半球』までちがった？　氷採掘は南極で合意、でも太陽光や有人は割れる。なぜ？　"
              "②電波天文台の班の場所から、地球は見える？（見えたら失格）　"
              "③赤道に基地を置くなら、ステップ1の1日約290Kの較差にどう対処する？　"
              "④このデータで『信じてよいか怪しいこと』は？（日照率・傾斜の絶対値／earth_elev_deg は秤動を無視／"
              "temp_amp_K は極で不確か）", 27)
    ws.column_dimensions["A"].width = 28
    for c in "BCDEF":
        ws.column_dimensions[c].width = 14

    # 並び順
    order = ["はじめに", "ステップ1_温度", "ステップ2_海と陸", S3G, S4R,
             "ステップ4b_南極", S4B, "ステップ5_まとめ", "データ_温度", "データ_クレーター",
             "データ_クレーター年代", "データ_環境", "データ_地域", "データ_南極", "データ_北極",
             "データ_地質", "データ_着陸地点", "参考"]
    wb._sheets.sort(key=lambda s: order.index(s.title) if s.title in order else 99)

    wb.save(OUT)
    print(f"wrote {OUT}  ({OUT.stat().st_size / 1024:.0f} KB, {len(wb.sheetnames)} シート)")


if __name__ == "__main__":
    build()
